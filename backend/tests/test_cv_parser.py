from io import BytesIO

import pytest
from docx import Document

from app.core.exceptions import (
    BadRequestError,
)
from app.services.recruitment.cv_parser import (
    CVParser,
)


@pytest.fixture
def parser():
    return CVParser()


def create_cv_docx() -> bytes:
    document = Document()

    document.add_paragraph(
        "Tinovimbanashe Shayamano"
    )
    document.add_paragraph(
        "tinovimba@example.com"
    )
    document.add_paragraph(
        "+263 77 123 4567"
    )
    document.add_heading(
        "Work Experience",
        level=1,
    )
    document.add_paragraph(
        "Data Analyst"
    )
    document.add_paragraph(
        "January 2024 – Present"
    )
    document.add_heading(
        "Education",
        level=1,
    )
    document.add_paragraph(
        "BSc Honours in Artificial Intelligence "
        "and Machine Learning"
    )
    document.add_heading(
        "Skills",
        level=1,
    )
    document.add_paragraph(
        "Python, SQL, Power BI, FastAPI and Docker"
    )

    stream = BytesIO()

    document.save(stream)

    return stream.getvalue()


def test_parser_extracts_candidate_details(
    parser,
):
    text = """
    Tinovimbanashe Shayamano
    tinovimba@example.com
    +263 77 123 4567

    Work Experience
    Data Analyst
    January 2024 – Present

    Education
    BSc Honours in Artificial Intelligence
    and Machine Learning

    Skills
    Python, SQL, Power BI, FastAPI, Docker
    """

    result = parser.parse(text)

    assert result["name"] == (
        "Tinovimbanashe Shayamano"
    )
    assert result["email"] == (
        "tinovimba@example.com"
    )
    assert result["phone"] == (
        "+263771234567"
    )
    assert result[
        "education_level"
    ] == "Bachelors"

    assert "python" in result[
        "extracted_skills"
    ]
    assert "sql" in result[
        "extracted_skills"
    ]
    assert "powerbi" in result[
        "extracted_skills"
    ]


def test_explicit_experience_is_preferred(
    parser,
):
    result = parser.parse(
        (
            "Jane Candidate\n"
            "Over 5 years of work experience "
            "in data analytics."
        )
    )

    assert result[
        "experience_years"
    ] == 5.0


def test_mojibake_dash_is_normalized(
    parser,
):
    result = parser.parse(
        (
            "Jane Candidate\n"
            "Work Experience\n"
            "Data Analyst\n"
            "2022 â€“ 2024\n"
            "Education\n"
        )
    )

    assert result[
        "experience_years"
    ] == 2.0


def test_overlapping_roles_are_not_double_counted(
    parser,
):
    result = parser.parse(
        """
        Jane Candidate

        Work Experience
        Data Analyst
        January 2022 - January 2024

        Reporting Analyst
        January 2023 - January 2024

        Education
        BSc Data Science
        """
    )

    assert result[
        "experience_years"
    ] == 2.0


def test_docx_text_is_extracted_and_parsed(
    parser,
):
    result = parser.parse_document(
        create_cv_docx(),
        ".docx",
    )

    assert result["name"] == (
        "Tinovimbanashe Shayamano"
    )
    assert result["email"] == (
        "tinovimba@example.com"
    )
    assert result[
        "education_level"
    ] == "Bachelors"
    assert "python" in result[
        "extracted_skills"
    ]


def test_corrupted_docx_is_rejected(
    parser,
):
    with pytest.raises(
        BadRequestError
    ):
        parser.parse_document(
            b"not a valid document",
            ".docx",
        )


def test_unsupported_document_type_is_rejected(
    parser,
):
    with pytest.raises(
        BadRequestError
    ):
        parser.parse_document(
            b"plain text",
            ".txt",
        )