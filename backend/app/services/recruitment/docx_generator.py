from io import BytesIO
from typing import Dict, Any, List
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class RecruitmentDocxGenerator:
    def generate_cv_docx(self, built_resume: Dict[str, Any]) -> BytesIO:
        doc = Document()
        self._setup_document(doc)
        self._add_cv_body(doc, built_resume)

        output = BytesIO()
        doc.save(output)
        output.seek(0)
        return output

    def generate_application_pack_docx(self, built_resume: Dict[str, Any]) -> BytesIO:
        doc = Document()
        self._setup_document(doc)

        self._add_cover_letter(doc, built_resume)
        doc.add_section(WD_SECTION.NEW_PAGE)
        self._add_cv_body(doc, built_resume)

        output = BytesIO()
        doc.save(output)
        output.seek(0)
        return output

    def _add_cv_body(self, doc: Document, built_resume: Dict[str, Any]):
        self._add_header(doc, built_resume)

        self._add_section_heading(doc, "Professional Summary")
        self._add_paragraph(doc, built_resume.get("professional_summary", ""))

        achievements = self._dedupe(built_resume.get("key_achievements", []))
        if achievements:
            self._add_section_heading(doc, "Key Achievements")
            self._add_bullets(doc, achievements)

        core_skills = self._dedupe(built_resume.get("core_skills", []))
        technical_skills = self._dedupe(built_resume.get("technical_skills", []))

        if core_skills:
            self._add_section_heading(doc, "Core Skills")
            self._add_skill_table(doc, core_skills)

        if technical_skills:
            self._add_section_heading(doc, "Technical Skills")
            self._add_skill_table(doc, technical_skills)

        self._add_section_heading(doc, "Professional Experience")
        self._add_experience(doc, built_resume.get("professional_experience", []))

        education = self._dedupe(built_resume.get("education", []))
        if education:
            self._add_section_heading(doc, "Education")
            self._add_bullets(doc, education)

        projects = self._dedupe(built_resume.get("projects", []))
        if projects:
            self._add_section_heading(doc, "Projects")
            self._add_bullets(doc, projects)

        certifications = self._dedupe(built_resume.get("certifications", []))
        if certifications:
            self._add_section_heading(doc, "Certifications")
            self._add_bullets(doc, certifications)

        languages = self._clean_languages(built_resume.get("languages", []))
        if languages:
            self._add_section_heading(doc, "Languages")
            self._add_bullets(doc, languages)

        self._add_section_heading(doc, "References")
        self._add_references(doc, built_resume.get("references"))

    def _setup_document(self, doc: Document):
        section = doc.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        self._set_default_font(doc)

    def _set_default_font(self, doc: Document):
        styles = doc.styles
        normal = styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10)

        for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
            style = styles[style_name]
            style.font.name = "Arial"
            style.font.color.rgb = RGBColor(15, 23, 42)

    def _add_header(self, doc: Document, built_resume: Dict[str, Any]):
        header = built_resume.get("header", {})
        name = header.get("name") or "Candidate"
        email = header.get("email") or ""
        phone = header.get("phone") or ""
        location = header.get("location") or ""

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(name.upper())
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(2, 132, 199)

        contact_parts = [x for x in [phone, email, location] if x]
        if contact_parts:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("  |  ".join(contact_parts))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(71, 85, 105)

        target = built_resume.get("target_role")
        organisation = built_resume.get("target_organisation")

        if target:
            target_text = f"Target Role: {target}"
            if organisation and organisation.lower() not in ["your organisation", "target organisation"]:
                target_text += f" | {organisation}"

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(target_text)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(51, 65, 85)

        self._add_horizontal_rule(doc)

    def _add_cover_letter(self, doc: Document, built_resume: Dict[str, Any]):
        header = built_resume.get("header", {})
        name = header.get("name") or "Candidate"
        email = header.get("email") or ""
        phone = header.get("phone") or ""
        location = header.get("location") or ""

        p = doc.add_paragraph()
        run = p.add_run(name.upper())
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(2, 132, 199)

        contact = " | ".join([x for x in [phone, email, location] if x])
        if contact:
            p = doc.add_paragraph()
            r = p.add_run(contact)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(71, 85, 105)

        self._add_horizontal_rule(doc)

        p = doc.add_paragraph()
        r = p.add_run("APPLICATION LETTER")
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(15, 23, 42)

        letter = built_resume.get("cover_letter") or ""
        for block in letter.split("\n\n"):
            clean = block.strip()
            if clean:
                self._add_paragraph(doc, clean)

    def _add_section_heading(self, doc: Document, text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)

        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(2, 132, 199)

        self._add_bottom_border(p)

    def _add_paragraph(self, doc: Document, text: str):
        if not text:
            return

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(str(text))
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(30, 41, 59)

    def _add_bullets(self, doc: Document, items: List[str]):
        for item in self._dedupe(items):
            p = doc.add_paragraph(style=None)
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            p.paragraph_format.space_after = Pt(2)

            run = p.add_run("• ")
            run.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199)

            r = p.add_run(str(item))
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(30, 41, 59)

    def _add_skill_table(self, doc: Document, skills: List[str]):
        skills = self._dedupe(skills)
        if not skills:
            return

        cols = 3
        rows = (len(skills) + cols - 1) // cols
        table = doc.add_table(rows=rows, cols=cols)
        table.autofit = True

        idx = 0
        for row in table.rows:
            for cell in row.cells:
                if idx < len(skills):
                    cell.text = str(skills[idx])
                    self._shade_cell(cell, "E0F2FE")
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            run.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    cell.text = ""
                idx += 1

    def _add_experience(self, doc: Document, experience: List[Dict[str, Any]]):
        if not experience:
            self._add_bullets(
                doc,
                [
                    "Supported accurate documentation, reporting, and operational tasks.",
                    "Maintained organised records and contributed to team objectives.",
                ],
            )
            return

        for item in experience:
            role = item.get("role", "Relevant Experience")
            organisation = item.get("organisation", "")
            period = item.get("period", "")

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)

            r = p.add_run(role)
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(15, 23, 42)

            if organisation:
                r = p.add_run(f" | {organisation}")
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(71, 85, 105)

            if period:
                r = p.add_run(f" | {period}")
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(100, 116, 139)

            self._add_bullets(doc, item.get("bullets", []))

    def _add_references(self, doc: Document, references):
        normalized_blocks = self._normalize_references_for_docx(references)

        if not normalized_blocks:
            self._add_paragraph(doc, "Available upon request.")
            return

        for index, block in enumerate(normalized_blocks):
            if index > 0:
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(2)

            for line_index, line in enumerate(block):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)

                r = p.add_run(line)
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)

                if line_index == 0:
                    r.bold = True

    def _normalize_references_for_docx(self, references) -> List[List[str]]:
        lines = self._reference_lines(references)
        lines = self._clean_reference_lines(lines)

        if not lines:
            return []

        blocks = self._group_reference_lines(lines)
        parsed_refs = [self._parse_reference_block(block) for block in blocks]

        organisation_hint = self._shared_reference_organisation(parsed_refs, lines)

        normalized = []
        seen_names = set()

        for ref in parsed_refs:
            if not ref.get("name"):
                continue

            if not ref.get("organisation") and organisation_hint:
                ref["organisation"] = organisation_hint

            block = self._render_reference_block(ref)

            if not block:
                continue

            name_key = self._normalise(block[0])
            if name_key in seen_names:
                continue

            seen_names.add(name_key)
            normalized.append(block)

        return normalized

    def _reference_lines(self, references) -> List[str]:
        if not references:
            return []

        if isinstance(references, str):
            return [line.strip() for line in references.splitlines() if line.strip()]

        if isinstance(references, list):
            return [str(line).strip() for line in references if str(line).strip()]

        return []

    def _clean_reference_lines(self, lines: List[str]) -> List[str]:
        output = []

        for line in lines:
            clean = str(line).strip()
            lower = clean.lower()

            if not clean:
                continue

            if lower in ["reference", "references", "professional references"]:
                continue

            if lower in ["english", "shona", "ndebele"]:
                continue

            if "interest" in lower or "football" in lower or "cricket" in lower:
                continue

            if clean not in output:
                output.append(clean)

        return output

    def _group_reference_lines(self, lines: List[str]) -> List[List[str]]:
        blocks = []
        current = []

        for line in lines:
            if self._looks_like_reference_name(line) and current:
                blocks.append(current)
                current = [line]
            else:
                current.append(line)

        if current:
            blocks.append(current)

        return blocks

    def _parse_reference_block(self, block: List[str]) -> Dict[str, str]:
        ref = {
            "name": "",
            "title": "",
            "organisation": "",
            "phone": "",
            "email": "",
        }

        for line in block:
            clean = str(line).strip()
            if not clean:
                continue

            if self._looks_like_phone(clean):
                ref["phone"] = self._format_phone(clean)
                continue

            if self._looks_like_email(clean):
                ref["email"] = self._format_email(clean)
                continue

            if self._looks_like_organisation(clean):
                if not ref["organisation"]:
                    ref["organisation"] = clean
                continue

            if self._looks_like_reference_title(clean):
                if not ref["title"]:
                    ref["title"] = clean
                continue

            if self._looks_like_reference_name(clean):
                if not ref["name"]:
                    ref["name"] = clean
                continue

            if not ref["title"]:
                ref["title"] = clean
            elif not ref["organisation"]:
                ref["organisation"] = clean

        return ref

    def _shared_reference_organisation(
        self,
        parsed_refs: List[Dict[str, str]],
        original_lines: List[str],
    ) -> str:
        orgs = []

        for ref in parsed_refs:
            org = ref.get("organisation")
            if org:
                orgs.append(org)

        if orgs:
            return max(set(orgs), key=orgs.count)

        for line in original_lines:
            if self._looks_like_organisation(line):
                return line

        return ""

    def _render_reference_block(self, ref: Dict[str, str]) -> List[str]:
        ordered = [
            ref.get("name", ""),
            ref.get("title", ""),
            ref.get("organisation", ""),
            ref.get("phone", ""),
            ref.get("email", ""),
        ]

        output = []
        seen = set()

        for line in ordered:
            clean = str(line).strip()
            key = self._normalise(clean)

            if not clean or key in seen:
                continue

            seen.add(key)
            output.append(clean)

        return output

    def _looks_like_reference_name(self, line: str) -> bool:
        clean = str(line).strip()

        if not clean:
            return False

        if ":" in clean:
            return False

        if self._looks_like_phone(clean):
            return False

        if self._looks_like_email(clean):
            return False

        if self._looks_like_organisation(clean):
            return False

        if self._looks_like_reference_title(clean):
            return False

        if any(char.isdigit() for char in clean):
            return False

        words = clean.split()

        if len(words) < 2 or len(words) > 4:
            return False

        first = words[0].lower().replace(".", "")
        if first in {"mr", "mrs", "ms", "dr", "prof"}:
            return True

        return all(word[0].isupper() or "." in word for word in words)

    def _looks_like_reference_title(self, line: str) -> bool:
        lower = str(line).lower().strip()

        titles = [
            "accountant",
            "chief accountant",
            "supervisor",
            "manager",
            "finance manager",
            "program coordinator",
            "programme coordinator",
            "officer",
            "director",
            "administrator",
            "lecturer",
            "chief",
            "head",
        ]

        return any(title in lower for title in titles) and not self._looks_like_organisation(line)

    def _looks_like_organisation(self, line: str) -> bool:
        lower = str(line).lower()

        org_words = [
            "ministry",
            "department",
            "council",
            "university",
            "college",
            "company",
            "agency",
            "authority",
            "organisation",
            "organization",
            "pvt",
            "ltd",
            "limited",
            "corporation",
            "infrastructure",
            "development",
            "transport",
            "environmental management",
            "ema",
        ]

        return any(word in lower for word in org_words)

    def _looks_like_phone(self, line: str) -> bool:
        lower = str(line).lower()
        digits = re.sub(r"\D", "", line)

        return (
            "phone" in lower
            or "cell" in lower
            or "mobile" in lower
            or str(line).strip().startswith("+")
            or len(digits) >= 9
        )

    def _format_phone(self, line: str) -> str:
        clean = str(line).strip()

        if clean.lower().startswith("phone"):
            return clean

        return f"Phone: {clean}"

    def _looks_like_email(self, line: str) -> bool:
        return bool(re.search(r"[\w\.-]+@[\w\.-]+\.\w+", str(line)))

    def _format_email(self, line: str) -> str:
        match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", str(line))
        if not match:
            return str(line).strip()

        return f"Email: {match.group(0).lower()}"

    def _normalise(self, text: str) -> str:
        value = str(text or "").lower()
        value = re.sub(r"[^a-z0-9]+", " ", value)
        value = re.sub(r"\s+", " ", value)
        return value.strip()

    def _clean_languages(self, languages: List[str]) -> List[str]:
        output = []

        for item in languages:
            text = str(item).strip()
            lower = text.lower()

            if not text:
                continue

            if any(bad in lower for bad in ["reference", "email:", "phone:", "mr ", "mrs ", "ms "]):
                continue

            if "interest" in lower or "football" in lower or "cricket" in lower:
                continue

            if text not in output:
                output.append(text)

        return output[:5]

    def _dedupe(self, items: List[str]) -> List[str]:
        output = []
        seen = set()

        for item in items or []:
            text = str(item).strip()
            if not text:
                continue

            normalized = text.lower().replace("  ", " ")
            if normalized in seen:
                continue

            seen.add(normalized)
            output.append(text)

        return output

    def _add_horizontal_rule(self, doc: Document):
        p = doc.add_paragraph()
        self._add_bottom_border(p)

    def _add_bottom_border(self, paragraph):
        p = paragraph._p
        p_pr = p.get_or_add_pPr()

        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)

        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "0284C7")
        p_bdr.append(bottom)

    def _shade_cell(self, cell, fill: str):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)


docx_generator = RecruitmentDocxGenerator()