"""
Recruitment CV parsing and document-text extraction.

Supported formats:
- PDF
- DOCX
"""

import re
import unicodedata
from datetime import date
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import (
    PackageNotFoundError,
)
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.config import settings
from app.core.exceptions import (
    BadRequestError,
)


class CVParser:
    SKILL_ALIASES = {
        "python": (
            "python",
        ),
        "sql": (
            "sql",
        ),
        "machine learning": (
            "machine learning",
            "ml engineering",
        ),
        "deep learning": (
            "deep learning",
        ),
        "nlp": (
            "nlp",
            "natural language processing",
        ),
        "react": (
            "react",
            "reactjs",
            "react.js",
        ),
        "fastapi": (
            "fastapi",
        ),
        "docker": (
            "docker",
            "containerisation",
            "containerization",
        ),
        "postgresql": (
            "postgresql",
            "postgres",
        ),
        "tableau": (
            "tableau",
        ),
        "excel": (
            "excel",
            "microsoft excel",
        ),
        "powerbi": (
            "power bi",
            "powerbi",
        ),
        "statistics": (
            "statistics",
            "statistical analysis",
        ),
        "data analysis": (
            "data analysis",
            "data analytics",
        ),
        "data visualization": (
            "data visualization",
            "data visualisation",
        ),
        "dashboard": (
            "dashboard",
            "dashboards",
        ),
        "pandas": (
            "pandas",
        ),
        "numpy": (
            "numpy",
        ),
        "matplotlib": (
            "matplotlib",
        ),
        "seaborn": (
            "seaborn",
        ),
        "scikit-learn": (
            "scikit-learn",
            "sklearn",
        ),
        "tensorflow": (
            "tensorflow",
        ),
        "keras": (
            "keras",
        ),
        "xgboost": (
            "xgboost",
        ),
        "lightgbm": (
            "lightgbm",
        ),
        "catboost": (
            "catboost",
        ),
        "r": (
            "r programming",
            "r language",
        ),
        "spss": (
            "spss",
        ),
        "stata": (
            "stata",
        ),
        "power query": (
            "power query",
        ),
        "dax": (
            "dax",
        ),
        "git": (
            "git",
        ),
        "github": (
            "github",
        ),
        "api": (
            "api",
            "rest api",
            "restful api",
        ),
        "etl": (
            "etl",
        ),
        "mysql": (
            "mysql",
        ),
        "data cleaning": (
            "data cleaning",
        ),
        "data capturing": (
            "data capturing",
            "data capture",
        ),
        "data verification": (
            "data verification",
            "data validation",
        ),
        "reporting": (
            "reporting",
            "report writing",
        ),
    }

    NAME_EXCLUDED_WORDS = {
        "curriculum",
        "vitae",
        "resume",
        "cv",
        "email",
        "phone",
        "mobile",
        "address",
        "profile",
        "summary",
        "objective",
        "experience",
        "education",
        "skills",
        "qualifications",
    }

    EXPERIENCE_HEADINGS = {
        "work experience",
        "professional experience",
        "employment history",
        "employment experience",
        "career history",
        "industrial attachment",
        "internship experience",
    }

    SECTION_STOP_HEADINGS = {
        "education",
        "academic background",
        "qualifications",
        "skills",
        "technical skills",
        "projects",
        "certifications",
        "references",
        "referees",
        "interests",
        "languages",
    }

    MONTHS = {
        "jan": 1,
        "january": 1,
        "feb": 2,
        "february": 2,
        "mar": 3,
        "march": 3,
        "apr": 4,
        "april": 4,
        "may": 5,
        "jun": 6,
        "june": 6,
        "jul": 7,
        "july": 7,
        "aug": 8,
        "august": 8,
        "sep": 9,
        "sept": 9,
        "september": 9,
        "oct": 10,
        "october": 10,
        "nov": 11,
        "november": 11,
        "dec": 12,
        "december": 12,
    }

    MONTH_PATTERN = (
        r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|"
        r"apr(?:il)?|may|jun(?:e)?|jul(?:y)?|"
        r"aug(?:ust)?|sep(?:t(?:ember)?)?|"
        r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
    )

    DATE_RANGE_PATTERN = re.compile(
        (
            rf"(?:(?P<start_month>{MONTH_PATTERN})"
            rf"\s+)?"
            rf"(?P<start_year>19\d{{2}}|20\d{{2}})"
            rf"\s*(?:-|–|—|to)\s*"
            rf"(?:(?P<end_month>{MONTH_PATTERN})"
            rf"\s+)?"
            rf"(?P<end_year>"
            rf"19\d{{2}}|20\d{{2}}|"
            rf"present|current|date|now)"
        ),
        flags=re.IGNORECASE,
    )

    def parse_document(
        self,
        content: bytes,
        extension: str,
    ) -> dict[str, Any]:
        raw_text = self.extract_text(
            content,
            extension,
        )

        return self.parse(
            raw_text
        )

    def extract_text(
        self,
        content: bytes,
        extension: str,
    ) -> str:
        normalized_extension = (
            extension.strip().lower()
        )

        if normalized_extension == ".pdf":
            raw_text = self._extract_pdf_text(
                content
            )

        elif normalized_extension == ".docx":
            raw_text = self._extract_docx_text(
                content
            )

        else:
            raise BadRequestError(
                (
                    "The CV document format cannot "
                    "be parsed."
                )
            )

        cleaned_text = self._normalise_text(
            raw_text
        )

        if not cleaned_text.strip():
            raise BadRequestError(
                (
                    "No readable text could be "
                    "extracted from the CV."
                ),
                details={
                    "reason": (
                        "empty_or_scanned_document"
                    ),
                },
            )

        if (
            len(cleaned_text)
            > settings.MAX_CV_TEXT_CHARACTERS
        ):
            raise BadRequestError(
                (
                    "The extracted CV text exceeds "
                    "the permitted processing size."
                )
            )

        return cleaned_text

    def parse(
        self,
        raw_text: str,
    ) -> dict[str, Any]:
        normalized_text = (
            self._normalise_text(
                raw_text or ""
            )
        )

        text_lower = (
            normalized_text.lower()
        )

        return {
            "name": self._extract_name(
                normalized_text
            ),
            "email": self._extract_email(
                normalized_text
            ),
            "phone": self._extract_phone(
                normalized_text
            ),
            "education_level": (
                self._extract_education(
                    text_lower
                )
            ),
            "experience_years": (
                self._extract_experience(
                    text_lower
                )
            ),
            "extracted_skills": (
                self._extract_skills(
                    text_lower
                )
            ),
        }

    def _extract_pdf_text(
        self,
        content: bytes,
    ) -> str:
        try:
            reader = PdfReader(
                BytesIO(content),
                strict=False,
            )

            if reader.is_encrypted:
                try:
                    result = reader.decrypt(
                        ""
                    )

                except Exception as exc:
                    raise BadRequestError(
                        (
                            "Password-protected PDF CVs "
                            "are not supported."
                        )
                    ) from exc

                if not result:
                    raise BadRequestError(
                        (
                            "Password-protected PDF CVs "
                            "are not supported."
                        )
                    )

            if (
                len(reader.pages)
                > settings.MAX_CV_PAGES
            ):
                raise BadRequestError(
                    (
                        "The PDF CV contains too many "
                        "pages."
                    ),
                    details={
                        "maximum_pages": (
                            settings.MAX_CV_PAGES
                        ),
                    },
                )

            extracted_pages: list[str] = []
            current_length = 0

            for page in reader.pages:
                page_text = (
                    page.extract_text() or ""
                )

                current_length += len(
                    page_text
                )

                if (
                    current_length
                    > settings.MAX_CV_TEXT_CHARACTERS
                ):
                    raise BadRequestError(
                        (
                            "The extracted PDF text "
                            "exceeds the permitted size."
                        )
                    )

                extracted_pages.append(
                    page_text
                )

            return "\n".join(
                extracted_pages
            )

        except BadRequestError:
            raise

        except (
            PdfReadError,
            ValueError,
            TypeError,
            OSError,
        ) as exc:
            raise BadRequestError(
                (
                    "The PDF CV is corrupted or "
                    "cannot be read."
                ),
                details={
                    "reason": (
                        "invalid_pdf_document"
                    ),
                },
            ) from exc

    def _extract_docx_text(
        self,
        content: bytes,
    ) -> str:
        try:
            document = Document(
                BytesIO(content)
            )

            sections: list[str] = []

            for paragraph in (
                document.paragraphs
            ):
                if paragraph.text.strip():
                    sections.append(
                        paragraph.text
                    )

            for table in document.tables:
                for row in table.rows:
                    row_values = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]

                    if row_values:
                        sections.append(
                            " | ".join(
                                row_values
                            )
                        )

            for section in document.sections:
                for paragraph in (
                    section.header.paragraphs
                ):
                    if paragraph.text.strip():
                        sections.append(
                            paragraph.text
                        )

                for paragraph in (
                    section.footer.paragraphs
                ):
                    if paragraph.text.strip():
                        sections.append(
                            paragraph.text
                        )

            return "\n".join(
                sections
            )

        except (
            BadZipFile,
            PackageNotFoundError,
            KeyError,
            ValueError,
            TypeError,
            OSError,
        ) as exc:
            raise BadRequestError(
                (
                    "The DOCX CV is corrupted or "
                    "cannot be read."
                ),
                details={
                    "reason": (
                        "invalid_docx_document"
                    ),
                },
            ) from exc

    def _normalise_text(
        self,
        text: str,
    ) -> str:
        normalized = unicodedata.normalize(
            "NFKC",
            text,
        )

        replacements = {
            "â€“": "–",
            "â€”": "—",
            "\u00a0": " ",
            "\x00": "",
        }

        for old, new in replacements.items():
            normalized = normalized.replace(
                old,
                new,
            )

        normalized = re.sub(
            r"[ \t]+",
            " ",
            normalized,
        )

        normalized = re.sub(
            r"\n{3,}",
            "\n\n",
            normalized,
        )

        return normalized.strip()

    def _extract_name(
        self,
        text: str,
    ) -> str | None:
        lines = [
            line.strip(" |,-")
            for line in text.splitlines()
            if line.strip()
        ]

        for line in lines[:12]:
            lowered = line.lower()

            if (
                "@" in line
                or "http://" in lowered
                or "https://" in lowered
                or "www." in lowered
            ):
                continue

            if re.search(
                r"\d",
                line,
            ):
                continue

            if any(
                word in lowered
                for word in self.NAME_EXCLUDED_WORDS
            ):
                continue

            words = line.split()

            if not 2 <= len(words) <= 5:
                continue

            if len(line) > 80:
                continue

            valid_words = all(
                re.fullmatch(
                    r"[A-Za-zÀ-ÖØ-öø-ÿ'’.-]+",
                    word,
                )
                is not None
                for word in words
            )

            if valid_words:
                return line.title()

        return None

    def _extract_email(
        self,
        text: str,
    ) -> str | None:
        match = re.search(
            (
                r"[a-zA-Z0-9._%+\-]+"
                r"@[a-zA-Z0-9.\-]+"
                r"\.[a-zA-Z]{2,}"
            ),
            text,
        )

        if not match:
            return None

        return (
            match.group(0)
            .strip(".,;:")
            .lower()
        )

    def _extract_phone(
        self,
        text: str,
    ) -> str | None:
        candidates = re.findall(
            r"\+?\d[\d\s().\-]{7,}\d",
            text,
        )

        for candidate in candidates:
            digits = re.sub(
                r"\D",
                "",
                candidate,
            )

            if not 8 <= len(digits) <= 15:
                continue

            if candidate.strip().startswith(
                "+"
            ):
                return f"+{digits}"

            return digits

        return None

    def _extract_skills(
        self,
        text: str,
    ) -> list[str]:
        found: set[str] = set()

        for canonical, aliases in (
            self.SKILL_ALIASES.items()
        ):
            for alias in aliases:
                pattern = (
                    r"(?<![A-Za-z0-9_])"
                    + re.escape(alias)
                    + r"(?![A-Za-z0-9_])"
                )

                if re.search(
                    pattern,
                    text,
                    flags=re.IGNORECASE,
                ):
                    found.add(
                        canonical
                    )

                    break

        return sorted(
            found
        )

    def _extract_education(
        self,
        text: str,
    ) -> str:
        education_patterns = [
            (
                "PhD",
                (
                    r"\bph\.?d\b",
                    r"\bdoctorate\b",
                    r"\bdoctoral\b",
                ),
            ),
            (
                "Masters",
                (
                    r"\bmaster(?:'s)?\b",
                    r"\bm\.?sc\b",
                    r"\bmba\b",
                    r"\bmcom\b",
                ),
            ),
            (
                "Bachelors",
                (
                    r"\bbachelor(?:'s)?\b",
                    r"\bb\.?sc\b",
                    r"\bbeng\b",
                    r"\bbcom\b",
                    r"\bhonou?rs\b",
                    r"\bhons\b",
                ),
            ),
            (
                "Diploma",
                (
                    r"\bdiploma\b",
                    r"\bhigher national diploma\b",
                    r"\bhnd\b",
                ),
            ),
            (
                "A Level",
                (
                    r"\ba[\s-]?level\b",
                    r"\badvanced level\b",
                ),
            ),
        ]

        for level, patterns in (
            education_patterns
        ):
            if any(
                re.search(
                    pattern,
                    text,
                )
                for pattern in patterns
            ):
                return level

        return "Unknown"

    def _extract_experience_section(
        self,
        text: str,
    ) -> str | None:
        lines = [
            line.strip()
            for line in text.splitlines()
        ]

        captured: list[str] = []
        active = False

        for line in lines:
            normalized_heading = (
                re.sub(
                    r"[^a-z ]",
                    "",
                    line.lower(),
                )
                .strip()
            )

            if (
                normalized_heading
                in self.EXPERIENCE_HEADINGS
            ):
                active = True
                continue

            if (
                active
                and normalized_heading
                in self.SECTION_STOP_HEADINGS
            ):
                break

            if active and line:
                captured.append(
                    line
                )

        if not captured:
            return None

        return "\n".join(
            captured
        )

    def _extract_experience(
        self,
        text: str,
    ) -> float | None:
        explicit_year_patterns = [
            (
                r"(\d+(?:\.\d+)?)\+?\s*years?"
                r"\s*(?:of\s*)?(?:work\s*)?experience"
            ),
            (
                r"experience[:\s]+"
                r"(\d+(?:\.\d+)?)\+?\s*years?"
            ),
            (
                r"(?:over|more than|at least)\s*"
                r"(\d+(?:\.\d+)?)\s*years?"
                r"\s*(?:of\s*)?(?:work\s*)?experience"
            ),
        ]

        for pattern in (
            explicit_year_patterns
        ):
            match = re.search(
                pattern,
                text,
            )

            if match:
                value = float(
                    match.group(1)
                )

                return round(
                    min(value, 40),
                    1,
                )

        month_match = re.search(
            (
                r"(\d+(?:\.\d+)?)\s*months?"
                r"\s*(?:of\s*)?(?:work\s*)?experience"
            ),
            text,
        )

        if month_match:
            months = float(
                month_match.group(1)
            )

            return round(
                min(
                    months / 12,
                    40,
                ),
                1,
            )

        experience_text = (
            self._extract_experience_section(
                text
            )
            or text
        )

        today = date.today()

        current_month_index = (
            today.year * 12
            + today.month
            - 1
        )

        intervals: list[
            tuple[int, int]
        ] = []

        for match in (
            self.DATE_RANGE_PATTERN.finditer(
                experience_text
            )
        ):
            start_year = int(
                match.group(
                    "start_year"
                )
            )

            start_month_name = (
                match.group(
                    "start_month"
                )
            )

            start_month = (
                self.MONTHS.get(
                    start_month_name.lower(),
                    1,
                )
                if start_month_name
                else 1
            )

            start_index = (
                start_year * 12
                + start_month
                - 1
            )

            if (
                start_index
                > current_month_index
            ):
                continue

            end_value = (
                match.group(
                    "end_year"
                )
                .lower()
            )

            end_month_name = (
                match.group(
                    "end_month"
                )
            )

            if end_value in {
                "present",
                "current",
                "date",
                "now",
            }:
                end_index = (
                    current_month_index
                )

            else:
                end_year = int(
                    end_value
                )

                if end_month_name:
                    end_month = (
                        self.MONTHS[
                            end_month_name.lower()
                        ]
                    )

                    end_index = (
                        end_year * 12
                        + end_month
                        - 1
                    )

                else:
                    # A year-only end value is treated as
                    # the beginning of that year.
                    end_index = (
                        end_year * 12
                    )

            end_index = min(
                end_index,
                current_month_index,
            )

            if end_index <= start_index:
                continue

            intervals.append(
                (
                    start_index,
                    end_index,
                )
            )

        if not intervals:
            return None

        intervals.sort()

        merged: list[
            list[int]
        ] = []

        for start, end in intervals:
            if (
                not merged
                or start
                > merged[-1][1]
            ):
                merged.append(
                    [
                        start,
                        end,
                    ]
                )

            else:
                merged[-1][1] = max(
                    merged[-1][1],
                    end,
                )

        total_months = sum(
            end - start
            for start, end in merged
        )

        if total_months <= 0:
            return None

        total_years = (
            total_months / 12
        )

        return round(
            min(total_years, 40),
            1,
        )


cv_parser = CVParser()